from pathlib import Path
from pptx import Presentation
import json


# Project directories
BASE_DIR = Path(__file__).resolve().parent.parent

RAW_DIR = BASE_DIR / "data" / "raw"
PROCESSED_DIR = BASE_DIR / "data" / "processed"

# Create processed folder if it doesn't exist
PROCESSED_DIR.mkdir(parents=True, exist_ok=True)


def extract_text_from_pptx(file_path):

    presentation = Presentation(file_path)

    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        combined_text = "\n".join(slide_text).strip()

        if combined_text:

            document = {
                "text": combined_text,

                "metadata": {
                    "source": file_path.name,
                    "file_type": "pptx",
                    "slide": slide_number,
                    "document_type": "Data Mining Unit 1"
                }
            }

            documents.append(document)

    return documents


def process_all_ppts():

    all_documents = []

    ppt_files = list(RAW_DIR.glob("*.pptx"))

    print(f"Found {len(ppt_files)} PPTX files.")

    for ppt_file in ppt_files:

        print(f"\nProcessing: {ppt_file.name}")

        documents = extract_text_from_pptx(ppt_file)

        print(f"Extracted {len(documents)} slides")

        all_documents.extend(documents)

    return all_documents


def save_documents(documents):

    output_file = PROCESSED_DIR / "documents.jsonl"

    with open(
        output_file,
        "w",
        encoding="utf-8"
    ) as f:

        for document in documents:

            f.write(
                json.dumps(
                    document,
                    ensure_ascii=False
                )
                + "\n"
            )

    print(f"\nSaved {len(documents)} documents to:")
    print(output_file)


if __name__ == "__main__":

    documents = process_all_ppts()

    save_documents(documents)

    print("\nIngestion completed successfully.")
